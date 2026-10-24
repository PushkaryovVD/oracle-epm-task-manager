"""Excel report of the schedule preview and, after a run, of the deployment results.

Pure openpyxl: no Tk objects, safe to call from a worker thread. Never receives or writes secrets.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import PlannedSchedule, ScheduleStatus, ValidationIssue
from app.services.deployment_ledger import LedgerEntry
from app.services.naming_service import MONTH_NAMES

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F5FAD")
_STATUS_FILL = {
    ScheduleStatus.READY: "E2F3E9",
    ScheduleStatus.COMPLETED: "E2F3E9",
    ScheduleStatus.READY_UNCHECKED: "FFF3D6",
    ScheduleStatus.SKIPPED_EXISTS: "E9EDF2",
    ScheduleStatus.SKIPPED: "E9EDF2",
    ScheduleStatus.BLOCKED: "FCE6E4",
    ScheduleStatus.FAILED: "FCE6E4",
    ScheduleStatus.UNKNOWN: "FDE2C8",
    ScheduleStatus.PENDING: "F3F5F8",
    ScheduleStatus.DEPLOYING: "E3EEFA",
    ScheduleStatus.VALIDATING: "E3EEFA",
}


@dataclass
class DeploymentRunInfo:
    """Outcome of one deployment run, for the report's Summary sheet."""

    started: datetime
    finished: datetime | None
    requested: int = 0
    completed: int = 0
    skipped: int = 0
    failed: int = 0
    unknown: int = 0
    not_started: int = 0
    blocked: int = 0
    stop_reason: str = ""
    stop_message: str = ""


@dataclass
class ReportData:
    rows: list[PlannedSchedule]
    summary: list[tuple[str, str]]                       # (label, value) pairs for the Summary sheet
    plan_errors: list[str] = field(default_factory=list)
    issues: list[ValidationIssue] = field(default_factory=list)
    unknown_columns: list[str] = field(default_factory=list)
    run: DeploymentRunInfo | None = None                 # None = preview only, nothing deployed
    job_ids: dict[str, str] = field(default_factory=dict)  # schedule name -> Oracle job id (from the ledger)
    history: list[LedgerEntry] = field(default_factory=list)


def _sheet(wb, title: str, headers: list[str], rows: list[list], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _fmt_time(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value else "(still running)"


def _duration(run: DeploymentRunInfo) -> str:
    if not run.finished:
        return ""
    seconds = int((run.finished - run.started).total_seconds())
    minutes, seconds = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours}:{minutes:02d}:{seconds:02d}"


def write_report_xlsx(path: str | Path, data: ReportData) -> Path:
    path = Path(path)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- Summary
    ws = wb.create_sheet("Summary")
    ws.append(["Schedule deployment report" if data.run else "Schedule deployment preview"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    for label, value in data.summary:
        ws.append([label, value])
    ws.append([])
    run = data.run
    if run is None:
        ws.append(["Note", "Preview only. Nothing was deployed and no changes were made in Oracle."])
    else:
        ws.append(["Deployment"])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        for label, value in (
            ("Run started", _fmt_time(run.started)),
            ("Run finished", _fmt_time(run.finished)),
            ("Duration", _duration(run)),
            ("Total schedules requested", run.requested),
            ("Successfully created", run.completed),
            ("Existing schedules skipped", run.skipped),
            ("Failed", run.failed),
            ("Unknown - check Oracle", run.unknown),
            ("Not started", run.not_started),
            ("Blocked (not deployed)", run.blocked),
            ("Stop reason", run.stop_reason or "(none - all eligible schedules processed)"),
            ("Stop message", run.stop_message),
        ):
            ws.append([label, value])
    if data.plan_errors:
        ws.append([])
        ws.append(["Errors"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for message in data.plan_errors:
            ws.append(["", message])
    for row in ws.iter_rows(min_row=2, max_col=1):
        if row[0].font is None or not row[0].font.b:
            row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")

    # ---- Schedules
    rows = [[
        r.index, r.name, f"{MONTH_NAMES[r.month.month - 1]} {r.month.year}", r.period_member, r.year_member,
        r.day_zero, r.status.value, data.job_ids.get(r.name, ""), " ".join(r.messages),
    ] for r in data.rows]
    ws = _sheet(wb, "Schedules" if run else "Preview",
                ["#", "Schedule Name", "Month", "Oracle Period", "Oracle Year", "Day Zero", "Status",
                 "Oracle Job ID", "Messages"], rows, [6, 42, 18, 14, 12, 12, 30, 14, 90])
    status_fill = {s.value: c for s, c in _STATUS_FILL.items()}
    for r_idx in range(2, ws.max_row + 1):
        ws.cell(r_idx, 6).number_format = "yyyy-mm-dd"
        ws.cell(r_idx, 9).alignment = Alignment(wrap_text=True, vertical="top")
        fill = status_fill.get(ws.cell(r_idx, 7).value)
        if fill:
            ws.cell(r_idx, 7).fill = PatternFill("solid", fgColor=fill)

    # ---- Validation issues / unknown columns
    _sheet(wb, "Validation Issues", ["Severity", "Location", "Message"],
           [[i.severity.value, i.location, i.message] for i in data.issues], [12, 40, 100])
    _sheet(wb, "Unknown Columns", ["Column header (not sent to Oracle)"],
           [[c] for c in data.unknown_columns], [50])

    # ---- Deployment history (local ledger of this environment)
    if data.history:
        _sheet(wb, "Deployment History",
               ["Schedule", "Status", "Year", "Period", "Day Zero", "Org Unit", "Template", "Job ID", "Attempts",
                "Updated", "Details"],
               [[e.schedule_name, e.status, e.year, e.period, e.day_zero, e.org_unit, e.template_name, e.job_id,
                 e.attempts, e.updated_at.replace("T", " "), e.details] for e in data.history],
               [40, 12, 10, 10, 12, 14, 30, 12, 9, 20, 80])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
